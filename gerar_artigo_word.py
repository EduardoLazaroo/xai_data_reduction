from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

OUTPUT_PATH = r"c:\Users\eduar\projetos\xai_data_reduction\artigo_xai_reduction.docx"

def add_heading(doc, text, level=1):
    doc.add_heading(text, level=level)


def add_paragraph(doc, text, justify=True, bold=False, size=11):
    p = doc.add_paragraph()
    if justify:
        p.alignment = WD_PARAGRAPH_ALIGNMENT.JUSTIFY
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = 'Calibri'


def build_document():
    doc = Document()

    p = doc.add_paragraph()
    p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    run = p.add_run('XAI aplicada à redução de dados em Machine Learning')
    run.bold = True
    run.font.size = Pt(18)
    run.font.name = 'Calibri'

    doc.add_paragraph('')

    add_heading(doc, '1. Introdução', 1)
    add_paragraph(doc, 'O avanço recente em Machine Learning tem levado ao desenvolvimento de modelos cada vez mais sofisticados e com alto poder preditivo. No entanto, esse aumento de performance muitas vezes é acompanhado por maior complexidade, o que dificulta a interpretação das decisões tomadas. Esse problema é especialmente relevante em áreas sensíveis, como saúde, finanças, educação e segurança, em que a explicabilidade das previsões é tão importante quanto a acurácia.')
    add_paragraph(doc, 'Nesse cenário, a área de XAI (Explainable Artificial Intelligence) surge como resposta para tornar os modelos mais transparentes e compreensíveis para seres humanos. O objetivo de XAI não é apenas responder “qual foi a previsão?”, mas também “por que a previsão foi feita dessa forma?”. Isso permite que especialistas entendam melhor o comportamento dos modelos, identifiquem falhas e interpretem decisões em contextos reais.')
    add_paragraph(doc, 'No entanto, a explicabilidade raramente é usada apenas como uma ferramenta final de análise. Há uma possibilidade muito interessante de usá-la também como parte do próprio processo de construção do modelo. Em particular, técnicas de explicabilidade podem ser empregadas para identificar quais variáveis são mais relevantes para a tarefa e, a partir disso, reduzir a dimensionalidade dos dados. Essa estratégia é especialmente útil em problemas com grande número de atributos, em que nem todas as variáveis têm igual contribuição para a predição.')
    add_paragraph(doc, 'A redução de dados é uma preocupação central em Machine Learning, pois dados de alta dimensionalidade podem gerar custo computacional mais elevado, maior risco de overfitting e menor interpretabilidade. Assim, a pergunta central desta pesquisa é: é possível usar explicações do modelo para reduzir o conjunto de atributos sem comprometer substancialmente o desempenho preditivo?')
    add_paragraph(doc, 'Este estudo busca responder essa questão por meio de um pipeline experimental em que um modelo baseline é treinado, suas decisões são explicadas com SHAP e LIME, e rankings das famílias filter, wrapper e embedded são comparados à seleção guiada por SHAP. O LIME é reservado à auditoria local de instâncias individuais; não é usado como seletor global.')

    add_heading(doc, '2. Contextualização do problema', 1)
    add_paragraph(doc, 'A redução de dimensionalidade é um problema clássico em Machine Learning. Em muitos datasets, especialmente em áreas como saúde ou biomedicina, é comum encontrar um grande número de variáveis, algumas das quais podem ser redundantes, pouco informativas ou até mesmo ruidosas. Esse fenômeno é especialmente importante porque, quando se aumenta o número de atributos, o modelo pode se tornar mais caro computacionalmente, menos robusto e menos interpretável.')
    add_paragraph(doc, 'A solução mais comum para esse problema é usar técnicas de seleção de atributos antes do treinamento do modelo. Métodos filter avaliam atributos por evidências univariadas, métodos wrapper re-treinam um estimador em diferentes subconjuntos e métodos embedded selecionam durante o ajuste. Neste estudo, mutual information representa filter, RFE representa wrapper e regressão logística L1 representa embedded.')
    add_paragraph(doc, 'A XAI oferece uma leitura do impacto de cada atributo no comportamento do modelo. O SHAP produz o ranking global usado na seleção guiada por XAI, enquanto o LIME explica decisões individuais próximas ao limiar e permite examinar a esparsidade da explicação local. Nenhum desses valores é interpretado como causalidade.')
    add_paragraph(doc, 'Essa mudança de perspectiva separa seleção e auditoria: a seleção é comparada com três famílias de baseline e a explicabilidade é avaliada por proxies operacionais declarados, complementados pela leitura local do LIME.')

    add_heading(doc, '3. Problema de pesquisa', 1)
    add_paragraph(doc, 'O problema central desta pesquisa pode ser formulado da seguinte maneira: é possível utilizar técnicas de explicabilidade para reduzir o número de atributos de um conjunto de dados sem degradar significativamente o desempenho do modelo preditivo?')
    add_paragraph(doc, 'Essa pergunta é relevante porque conecta duas áreas que, em muitos projetos, são tratadas separadamente: Machine Learning, focado em performance, e XAI, focado em interpretabilidade. O objetivo do estudo é justamente unir esses dois aspectos em um único fluxo de pesquisa, em que a explicabilidade não serve apenas para justificar o modelo, mas também para melhorar sua representação interna.')

    add_heading(doc, '4. Objetivos', 1)
    add_paragraph(doc, 'Objetivo geral: Investigar se a seleção guiada pelo ranking global do SHAP reduz a dimensionalidade sem comprometer significativamente o desempenho, quando comparada a métodos filter, wrapper e embedded.')
    add_paragraph(doc, 'Objetivos específicos: (i) treinar um modelo baseline; (ii) gerar rankings por mutual information, RFE e regressão logística L1; (iii) gerar um ranking global por SHAP; (iv) reservar o LIME para auditoria de instâncias locais; (v) comparar desempenho, custo, redução dimensional e proxies de interpretabilidade; (vi) re-treinar o modelo reduzido e registrar limitações e referências atuais.')

    add_heading(doc, '5. Hipóteses', 1)
    add_paragraph(doc, 'H1: A utilização de técnicas de explicabilidade para orientar a redução de atributos permite diminuir a dimensionalidade do conjunto de dados sem prejudicar de forma significativa o desempenho do modelo.')
    add_paragraph(doc, 'H2: A seleção guiada por SHAP produz desempenho e custo competitivos em relação a representantes filter, wrapper e embedded, além de oferecer uma trilha explicativa global complementada pela auditoria local do LIME.')
    add_paragraph(doc, 'H3: A redução pode aumentar a concentração da importância global e a esparsidade das explicações locais, medidas como proxies operacionais e não como prova de causalidade.')

    add_heading(doc, '6. Justificativa', 1)
    add_paragraph(doc, 'A justificativa da pesquisa está na necessidade de combinar performance e interpretabilidade em um contexto de dados de alta dimensionalidade. Em muitos problemas reais, o maior desafio não é apenas classificar corretamente, mas também entender o que está sendo usado para decidir. Isso é especialmente importante em áreas em que a decisão final precisa ser explicável para especialistas humanos.')
    add_paragraph(doc, 'Além disso, a redução de atributos pode trazer benefícios práticos relevantes, como menor custo computacional, maior agilidade no treinamento, menor uso de memória, menor tempo de inferência, modelos mais simples e maior interpretabilidade. Dessa forma, a pesquisa tem relevância tanto para a área acadêmica quanto para aplicações práticas. Ela investiga se a explicabilidade pode ser usada como mecanismo de seleção e otimização do próprio pipeline de dados.')

    add_heading(doc, '7. Fundamentação teórica', 1)
    add_paragraph(doc, '7.1. Machine Learning e alta dimensionalidade: dados de alta dimensionalidade são comuns em aplicações de diagnóstico, sensores, dados genéticos, dados clínicos e sistemas de monitoramento. Nesses cenários, o número de atributos pode ser grande demais para a quantidade de observações disponível, o que aumenta a chance de overfitting e reduz a robustez do modelo.')
    add_paragraph(doc, '7.2. Seleção de atributos: a seleção de atributos é uma das formas mais comuns de lidar com esse problema. O objetivo é escolher um subconjunto de variáveis que seja suficiente para representar o problema, reduzindo ruído e redundância. Métodos filter, wrapper e embedded fazem escolhas com diferentes custos e dependências do estimador; neste projeto, mutual information, RFE e regressão logística L1 representam essas famílias.')
    add_paragraph(doc, '7.3. SHAP: SHAP (Shapley Additive Explanations) é uma técnica de explicabilidade baseada na teoria dos jogos cooperativos. Ela atribui valores de importância a cada atributo com base na contribuição marginal da variável para a previsão final do modelo. A principal vantagem do SHAP é que ele oferece uma explicação global e local, permitindo identificar quais atributos mais influenciam a saída do modelo e em que intensidade.')
    add_paragraph(doc, '7.4. LIME: LIME (Local Interpretable Model-agnostic Explanations) é outra técnica de explicabilidade. Diferentemente do SHAP, o LIME busca explicar uma previsão específica por meio de perturbações locais do dado. Isso torna a explicação mais interpretável em casos de decisões individuais.')
    add_paragraph(doc, '7.5. XAI e avaliação da interpretabilidade: o SHAP é usado para orientar a seleção global, enquanto o LIME permanece como explicador local. A pesquisa mede concentração da massa de |SHAP| no Top-10, esparsidade da explicação local LIME e redução de atributos como proxies operacionais. Esses indicadores não medem causalidade e devem ser discutidos junto de fidelidade, estabilidade e avaliação humana.')

    add_heading(doc, '8. Metodologia', 1)
    add_paragraph(doc, 'A metodologia foi organizada em etapas, seguindo uma lógica de estudo experimental.')
    add_paragraph(doc, '8.1. Geração do conjunto sintético: o primeiro passo foi gerar um conjunto de dados sintéticos com características semelhantes às de um problema de saúde ou diagnóstico, incluindo atributos informativos, redundantes, ruidosos, desbalanceamento leve e estrutura controlada para permitir análise comparativa.')
    add_paragraph(doc, '8.2. Treinamento do modelo baseline: em seguida, um modelo de referência foi treinado com todas as variáveis. Esse modelo serviu como linha de base para avaliar todas as posteriores reduções.')
    add_paragraph(doc, '8.3. Métricas de avaliação: as métricas usadas foram acurácia, precisão, recall, F1-score, ROC-AUC, tempo de treinamento, tempo de inferência e número de atributos. A combinação dessas métricas permite avaliar não só a acurácia do modelo, mas também o custo computacional da solução.')
    add_paragraph(doc, '8.4. Explicabilidade global e local: depois do treinamento, o SHAP foi usado para gerar um ranking global e o LIME para analisar uma instância próxima do limiar de decisão. A saída local do LIME não foi usada para ordenar a população.')
    add_paragraph(doc, '8.5. Baselines de seleção: foram construídos rankings de mutual information (filter), RFE com Random Forest (wrapper) e regressão logística L1 (embedded). Cada ranking foi avaliado no estudo de ablação progressiva, junto do ranking SHAP.')
    add_paragraph(doc, '8.6. Avaliação de interpretabilidade: além do número de atributos, foram calculadas a concentração da massa de |SHAP| no Top-10 e a esparsidade da explicação local LIME. São proxies operacionais para esta demonstração e não substituem análise de estabilidade, fidelidade ou avaliação por especialistas.')
    add_paragraph(doc, '8.7. Filtro híbrido e shap-select: além da seleção direta por SHAP, a metodologia inclui um estágio adicional de pré-filtro híbrido, com remoção de variáveis de baixa variância e alta correlação, e uma etapa de shap-select, em que os coeficientes de regressão sobre valores SHAP são usados para decidir quais atributos devem permanecer no modelo.')
    add_paragraph(doc, '8.8. Otimização dos hiperparâmetros: após a redução dos atributos, o modelo foi re-otimizado por hyperparameter search, utilizando Optuna. Isso permite avaliar se a etapa final do pipeline pode ser ajustada de maneira mais eficiente em um espaço menor de atributos.')
    add_paragraph(doc, '8.9. Dashboard comparativo final: por fim, os resultados foram reunidos em um dashboard comparativo para mostrar a diferença entre o modelo baseline completo e o modelo reduzido, em termos de número de atributos, desempenho, custo e proxies de interpretabilidade.')

    add_heading(doc, '9. Resultados esperados', 1)
    add_paragraph(doc, 'A expectativa desta pesquisa é que a redução guiada por XAI promova um equilíbrio interessante entre desempenho e simplicidade. Em outras palavras, espera-se que um modelo treinado com um conjunto menor de atributos mantenha métricas de desempenho próximas às do modelo completo, ao mesmo tempo em que reduza dimensionalidade, custo computacional e tempo de execução.')
    add_paragraph(doc, 'Em um cenário ideal, a redução pode até melhorar o desempenho por reduzir ruído e redundância. O ponto mais relevante é que a performance não deve cair de forma significativa quando os atributos menos úteis forem removidos.')

    add_heading(doc, '10. Resultados observados', 1)
    add_paragraph(doc, 'Os experimentos executados corroboram essa expectativa no cenário proposto. O modelo baseline foi treinado com 40 atributos, enquanto a estratégia guiada por XAI reduziu para 10 atributos. O modelo reduzido manteve o F1-score em nível equivalente ou superior ao baseline, e a redução de atributos foi expressiva.')
    add_paragraph(doc, 'Dados observados na execução validada: baseline com 40 atributos; modelo reduzido com 10 atributos; F1-score baseline de 0,8373; F1-score reduzido de 0,8475; redução de dimensionalidade de 75,0%; concentração SHAP Top-10 de 0,6689 e esparsidade local LIME de 75,0%. Esses números são específicos do split controlado e não constituem garantia clínica ou superioridade universal.')

    add_heading(doc, '11. Discussão', 1)
    add_paragraph(doc, 'Os resultados sugerem que o ranking SHAP pode funcionar como mecanismo útil de seleção neste protocolo. A comparação com mutual information, RFE e logística L1 torna a conclusão mais defensável, pois permite observar se o comportamento é específico da XAI ou se também aparece em outras famílias. O LIME acrescenta uma leitura local, mas não transforma uma explicação individual em evidência populacional.')
    add_paragraph(doc, 'Essa conclusão é relevante porque mostra que a redução da dimensionalidade não precisa ser feita apenas por critério estatístico puro ou por abordagem baseada em tentativa e erro; também pode ser guiada pela explicação do modelo. Isso cria uma ponte entre dois paradigmas que frequentemente aparecem separados: a otimização preditiva e a interpretação do modelo.')
    add_paragraph(doc, 'No entanto, é preciso reconhecer que essa estratégia também apresenta limitações. O modelo depende diretamente da qualidade do treinamento inicial. Se o classificador base for instável, as explicações podem refletir vieses do modelo e não necessariamente a estrutura real do problema. Além disso, o custo de computação da explicabilidade pode ser elevado em datasets muito grandes.')
    add_paragraph(doc, 'Portanto, a abordagem é promissora, mas deve ser interpretada como um trade-off entre desempenho preditivo, explicabilidade, custo computacional, simplicidade do modelo e qualidade do conjunto de dados. Esse equilíbrio é central para a discussão final do trabalho.')

    add_heading(doc, '12. Limitações do estudo', 1)
    add_paragraph(doc, 'Apesar dos resultados promissores, o estudo apresenta algumas limitações importantes. Primeiro, o uso de dados sintéticos oferece controle e clareza, mas não substitui a complexidade de dados do mundo real. Segundo, a pesquisa se concentra principalmente na redução de atributos e não em outras formas de redução, como redução de instâncias ou compressão temporal. Terceiro, métodos como SHAP e otimização por Optuna exigem tempo e custo computacional. Quarto, a qualidade das explicações depende diretamente da qualidade do modelo treinado.')

    add_heading(doc, '13. Contribuição do estudo', 1)
    add_paragraph(doc, 'A principal contribuição desta pesquisa é demonstrar que explicabilidade e redução de dimensionalidade podem ser integradas em um mesmo pipeline. Em vez de tratar XAI apenas como ferramenta de interpretação final, o estudo mostra que ela pode ter um papel ativo na etapa de seleção de variáveis e na melhoria da eficiência do modelo.')
    add_paragraph(doc, 'Além disso, a comparação com filter, wrapper e embedded fortalece a pesquisa ao situar a explicabilidade no panorama metodológico completo. A contribuição está em mostrar como uma seleção guiada por SHAP pode ser auditada por LIME e comparada com baselines conhecidos, sempre separando desempenho preditivo de interpretabilidade.')

    add_heading(doc, '14. Considerações finais', 1)
    add_paragraph(doc, 'Este estudo mostra que a utilização de explicabilidade como mecanismo de redução de dados é uma linha de pesquisa promissora. O modelo reduzido manteve ou melhorou a performance em um cenário de alta dimensionalidade, sugerindo que parte do conjunto original de atributos era redundante ou pouco informativa.')
    add_paragraph(doc, 'A partir dessa evidência, conclui-se apenas que o uso de SHAP para seleção de atributos é viável e promissor dentro do protocolo sintético executado. A pesquisa também ressalta que interpretabilidade exige critérios explícitos, comparação com baselines, estabilidade e validação em dados externos.')
    add_paragraph(doc, 'Em termos práticos, o trabalho sugere que a explicabilidade pode ser usada como estratégia de engenharia de dados, e não apenas como uma camada analítica final. Essa mudança de perspectiva é uma contribuição importante para a área.')

    add_heading(doc, '15. Referências selecionadas', 1)
    add_paragraph(doc, 'RIBEIRO, M. T.; SINGH, S.; GUESTRIN, C. Why Should I Trust You? Explaining the Predictions of Any Classifier. KDD, 2016. https://arxiv.org/abs/1602.04938')
    add_paragraph(doc, 'LUNDBERG, S. M.; LEE, S.-I. A Unified Approach to Interpreting Model Predictions. NeurIPS, 2017. https://papers.nips.cc/paper/7062-a-unified-approach-to-interpreting-model-predictions')
    add_paragraph(doc, 'MOLNAR, C. Interpretable Machine Learning. 2nd ed., 2022. https://christophm.github.io/interpretable-ml-book/')
    add_paragraph(doc, 'NIST. Artificial Intelligence Risk Management Framework (AI RMF 1.0), 2023. https://www.nist.gov/itl/ai-risk-management-framework')
    add_paragraph(doc, 'Resumo: Este estudo compara uma seleção guiada pelo ranking global do SHAP com métodos filter, wrapper e embedded. O LIME é reservado à auditoria local. A avaliação combina desempenho, custo, redução dimensional e proxies explícitos de interpretabilidade, sem confundir contribuição do modelo com causalidade.', bold=False, size=11)

    return doc


doc = build_document()
doc.save(OUTPUT_PATH)
print(f'Arquivo Word criado em: {OUTPUT_PATH}')
